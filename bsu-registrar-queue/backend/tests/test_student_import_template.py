import io

import openpyxl

from app.services.student_import import reader, template
from app.services.student_import.columns import COLUMNS, LABELS, REQUIRED


def test_template_has_headers_example_and_dropdowns():
    data = template.build_template()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Students"]
    headers = [c.value for c in ws[1]]
    assert headers == [LABELS[c] + ("*" if c in REQUIRED else "") for c in COLUMNS]
    assert ws.cell(row=2, column=1).value == "2026000001"
    assert ws.cell(row=2, column=1).number_format == "@"
    sqrefs = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert "F2" in sqrefs and "H2" in sqrefs            # Course, Year Level dropdowns
    assert all(dv.errorStyle == "warning" for dv in ws.data_validations.dataValidation)
    assert "Instructions" in wb.sheetnames


def test_untouched_template_imports_as_zero_rows():
    assert reader.read_upload(template.build_template()).rows == []
